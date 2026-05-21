import io
from models.request import InputType, ParallaxRequest


def parse_content(request: ParallaxRequest) -> str:

    if request.input_type == InputType.TEXT:
        return request.content or ""

    if request.input_type == InputType.FILE:
        return _parse_file(request.file_name, request.file_data)

    if request.input_type == InputType.IMAGE:
        return request.file_name or ""

    if request.input_type == InputType.VIDEO:
        return request.file_name or ""

    return ""


def _parse_file(file_name: str, file_data: bytes) -> str:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext in ["txt", "md", "csv"]:
        return file_data.decode("utf-8", errors="ignore")

    if ext == "pdf":
        return _parse_pdf(file_data)

    if ext == "docx":
        return _parse_docx(file_data)

    if ext == "xlsx":
        return _parse_xlsx(file_data)

    return ""


def _parse_pdf(file_data: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return ""


def _parse_docx(file_data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_data))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


def _parse_xlsx(file_data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_data), read_only=True)
        texts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                texts.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(texts)
    except Exception:
        return ""